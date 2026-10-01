import os, sys, io
import docx
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = os.path.abspath('report/報告書_前模擬.docx')
OUT = os.path.abspath('report/報告書_前模擬_new.docx')

doc = docx.Document(SRC)

# Mapping of circuit diagram table keywords to image path and width
CIRCUIT_REPLACEMENTS = {
    'SOFT_START 子電路': ('diagram/report/fig2_6_softstart.png', 14.5),
    'EA 子電路': ('diagram/report/fig2_8_ea.png', 10.5),
    'PWM_OSC 子電路': ('diagram/report/fig2_9_pwm_osc.png', 14.5),
    'CMP01 子電路': ('diagram/report/fig2_10_cmp01.png', 15.0),
    'TRIANGO_HZ 子電路': ('diagram/report/fig2_12_ontime.png', 14.0),
    'MUX2 子電路': ('diagram/report/fig2_14_mux2.png', 15.0),
    'DEAD_ZONE 子電路': ('diagram/report/fig2_15_deadzone.png', 14.5),
}

# 1. Process Tables
tables_to_remove = []
for t in list(doc.tables):
    txt = ' '.join(c.text.strip().replace('\n', ' ') for r in t.rows for c in r.cells)
    
    # Check circuit replacement
    replaced = False
    for key, (img_path, width_cm) in CIRCUIT_REPLACEMENTS.items():
        if key in txt:
            # Create replacement paragraph with image
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = docx.shared.Pt(6)
            p.paragraph_format.space_after = docx.shared.Pt(6)
            run = p.add_run()
            abs_img = os.path.abspath(img_path)
            run.add_picture(abs_img, width=Cm(width_cm))
            
            t._tbl.addprevious(p._p)
            t._tbl.getparent().remove(t._tbl)
            print(f'Replaced circuit table [{key}] with image {img_path}')
            replaced = True
            break
            
    if replaced:
        continue
        
    # Check waveform table to delete
    if '【請貼上 WaveView 波形截圖】' in txt or 'WaveView 波形截圖' in txt:
        t._tbl.getparent().remove(t._tbl)
        print(f'Deleted waveform table: {txt[:60]}...')

# 2. Process Paragraphs: Captions, Deletions, and Text Updates
# Captions to delete
CAPTIONS_TO_DELETE = [
    '圖2.7 軟啟動模擬結果',
    '圖2.12 PWM控制迴路模擬結果',
    '圖2.15 PFM控制迴路模擬結果',
    '圖2.17 零電流偵測與二極體模擬模擬結果',
    '圖2.20 死區控制電路模擬結果',
    '圖2.21 切換節點SW波形',
    '圖3.1 整體模擬波形',
    '圖3.2 啟動過程模擬結果',
    '圖3.3 重載PWM穩態模擬結果',
    '圖3.4 輕載PFM穩態模擬結果',
    '圖3.5 Undershoot響應（50 mA → 500 mA）',
    '圖3.6 Overshoot響應（500 mA → 50 mA）',
    '圖3.7 模式切換模擬結果',
]

# Paragraphs to delete entirely
PARAS_TO_DELETE = [
    'PWM控制迴路的模擬結果如圖2.12所示，RAMP與VCOMP比較後產生的工作週期訊號經多工器輸出為VDUTY_SEL。',
    '切換節點SW的波形如圖2.21所示。',
]

# Captions to rename
CAPTION_RENAMES = {
    '圖2.8 PWM控制迴路方塊圖': '圖2.7 PWM控制迴路方塊圖',
    '圖2.9 誤差放大器與補償網路': '圖2.8 誤差放大器',
    '圖2.10 鋸齒波產生器': '圖2.9 鋸齒波產生器',
    '圖2.11 比較器電路（CMP01）': '圖2.10 比較器電路（CMP01）',
    '圖2.13 PFM控制迴路方塊圖': '圖2.11 PFM控制迴路方塊圖',
    '圖2.14 導通時間產生器': '圖2.12 導通時間產生器',
    '圖2.16 PWM/PFM模式偵測方塊圖': '圖2.13 PWM/PFM模式偵測方塊圖',
    '圖2.18 多工器電路': '圖2.14 多工器電路',
    '圖2.19 死區控制電路': '圖2.15 死區控制電路',
}

# In-text replacements for specific paragraphs
# Note: we will iterate through paragraphs
for p in list(doc.paragraphs):
    text = p.text.strip()
    
    # 1. Check deletion of captions
    if p.style.name == 'FigCaption' and text in CAPTIONS_TO_DELETE:
        print(f'Deleting FigCaption: {text}')
        p._p.getparent().remove(p._p)
        continue
        
    # 2. Check deletion of pure intro paragraphs
    if text in PARAS_TO_DELETE:
        print(f'Deleting intro paragraph: {text}')
        p._p.getparent().remove(p._p)
        continue
        
    # 3. Rename FigCaptions
    if p.style.name == 'FigCaption' and text in CAPTION_RENAMES:
        new_cap = CAPTION_RENAMES[text]
        print(f'Renaming caption: {text} -> {new_cap}')
        p.text = new_cap
        continue
        
    # 4. Text modifications for waveform data paragraphs
    if text.startswith('軟啟動模擬結果如圖2.7所示。'):
        p.text = ('模擬中，VREF在5 μs、10 μs、15 μs時分別為0.32 V、0.83 V、1.18 V，'
                  '約在13.9 μs達到1.14 V，最後穩定在1.2778 V，峰對峰漣波為2.55 mV。'
                  '輸出電壓隨VREF平穩上升，在22.5 μs達到1.782 V（最終值的99 %），'
                  '8～32 μs之間輸出電壓的最大值為1.8031 V，過衝量僅0.17 %。')
        print('Updated 軟啟動模擬結果 text')
        continue
        
    if text.startswith('PFM控制迴路的模擬結果如圖2.15所示。'):
        p.text = ('模擬中，輕載時在62 μs之後連續8個脈衝的週期為443.4 ns～452.4 ns，'
                  '70 μs之後為448.0 ns～452.3 ns。')
        print('Updated PFM控制迴路模擬結果 text')
        continue

    if text.startswith('零電流偵測與二極體模擬的模擬結果如圖2.17所示。'):
        p.text = ('模擬中，輕載時（62～71 μs）電感電流的範圍為−11.7 mA～149 mA，'
                  '反向電流被限制在約12 mA以內；PFM_IDLE在輕載時會拉到3.3 V，重載時則保持在0 V。')
        print('Updated 零電流偵測與二極體模擬 text')
        continue

    if text.startswith('死區控制電路的模擬結果如圖2.20所示，'):
        p.text = '模擬中，上臂與下臂的閘極訊號在切換時不會同時處於導通狀態。'
        print('Updated 死區控制電路模擬結果 text')
        continue

    # Chapter 3 updates
    if text.startswith('整體模擬波形如圖3.1所示，'):
        p.text = ('整體前模擬涵蓋輸出電壓的啟動過程、負載跳變時的暫態響應以及MODE訊號隨負載切換的情形，'
                  '各項模擬結果與波形分析分別於後續各節進行說明。')
        print('Updated 3-1 結尾 text')
        continue

    if text.startswith('啟動過程如圖3.2所示。'):
        p.text = ('模擬中，輸出電壓在22.5 μs時達到最終值的99 %（1.782 V），'
                  '並在28.1 μs時第一次到達1.8 V；啟動期間輸出電壓最大值為1.8031 V，過衝為0.17 %。'
                  'REF_READY在31.7 μs時成立，啟動期間MODE保持在PFM，沒有誤切換。')
        print('Updated 3-2 啟動過程 text')
        continue

    if text.startswith('重載500 mA的穩態波形如圖3.3所示。'):
        p.text = ('模擬中，42～51 μs之間重載500 mA的平均輸出電壓為1.7999 V，'
                  '最小值1.7975 V、最大值1.8022 V，輸出漣波為4.70 mV。'
                  '切換頻率為2.2157 MHz，MODE維持在3.27 V以上（PWM模式）。')
        print('Updated 3-3 重載PWM穩態 text')
        continue

    if text.startswith('輕載50 mA的穩態波形如圖3.4所示。'):
        p.text = ('模擬中，輕載50 mA在62～71 μs之間平均輸出電壓為1.8030 V，輸出漣波為3.32 mV；'
                  '在較長的62～79 μs區間內平均輸出電壓同樣為1.8030 V，漣波為3.42 mV，'
                  'MODE最高只有0.03 V（PFM模式）。')
        print('Updated 3-4 輕載PFM穩態 text')
        continue

    if text.startswith('負載由50 mA跳到500 mA時的響應如圖3.5所示。'):
        p.text = ('模擬中，負載由50 mA跳到500 mA時，跳變前（28～31.9 μs）輸出電壓在1.7975 V～1.8030 V之間，'
                  '跳變後輸出電壓最低降到1.7917 V，相對於1.8 V的下衝為8.3 mV（0.46 %）。')
        print('Updated 3-5 Undershoot text')
        continue

    if text.startswith('負載由500 mA回到50 mA時的響應如圖3.6所示，'):
        p.text = ('負載由500 mA回到50 mA時，輸出電壓最高升到1.8133 V，'
                  '相對於1.8 V的過衝為13.3 mV（0.74 %）；52～60 μs之間輸出電壓的擺幅為13.5 mV。')
        print('Updated 3-5 Overshoot text')
        continue

    if text.startswith('模式切換的過程如圖3.7所示。'):
        p.text = ('模擬中，負載在32 μs增加後，電感電流不再降到零，N_TMR開始充電，'
                  'MODE在34.03 μs升到高電位，PFM切換到PWM的延遲為2.04 μs。'
                  '負載在52 μs減少後，零電流事件重新出現使N_TMR放電，'
                  'MODE在53.14 μs降回低電位，PWM切換到PFM的延遲為1.14 μs。')
        print('Updated 3-6 模式切換 text')
        continue

    # In-text replacements for figure numbers and transistor names
    modified = False
    new_t = text
    
    if 'PWM控制迴路如圖2.8所示' in new_t:
        new_t = new_t.replace('PWM控制迴路如圖2.8所示', 'PWM控制迴路如圖2.7所示')
        modified = True
    if '誤差放大器如圖2.9所示' in new_t:
        new_t = new_t.replace('誤差放大器如圖2.9所示', '誤差放大器如圖2.8所示')
        modified = True
    if '如圖2.10所示' in new_t and '鋸齒波產生器' in new_t:
        new_t = new_t.replace('如圖2.10所示', '如圖2.9所示')
        modified = True
    if '如圖2.11所示的比較器CMP01' in new_t:
        new_t = new_t.replace('如圖2.11所示的比較器CMP01', '如圖2.10所示的比較器CMP01')
        modified = True
    if 'M_MN3與M_MN4' in new_t:
        new_t = new_t.replace('M_MN3與M_MN4', 'MN3與MN4')
        modified = True
    if 'PFM控制迴路如圖2.13所示' in new_t:
        new_t = new_t.replace('PFM控制迴路如圖2.13所示', 'PFM控制迴路如圖2.11所示')
        modified = True
    if '導通時間產生器如圖2.14所示' in new_t:
        new_t = new_t.replace('導通時間產生器如圖2.14所示', '導通時間產生器如圖2.12所示')
        modified = True
    if 'PWM/PFM模式偵測方塊如圖2.16所示' in new_t:
        new_t = new_t.replace('PWM/PFM模式偵測方塊如圖2.16所示', 'PWM/PFM模式偵測方塊如圖2.13所示')
        modified = True
    if '多工器如圖2.18所示' in new_t:
        new_t = new_t.replace('多工器如圖2.18所示', '多工器如圖2.14所示')
        modified = True
    if '如圖2.19所示' in new_t and '死區控制電路' in new_t:
        new_t = new_t.replace('如圖2.19所示', '如圖2.15所示')
        modified = True

    if modified:
        print(f'Modified in-text: {new_t[:60]}...')
        p.text = new_t

doc.save(OUT)
print('Document saved to:', OUT)
